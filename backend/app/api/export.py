from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from app.core.database import get_db
from app.models.user import User
from app.api.auth import get_current_user
from app.crud import villager, household, contact, bank_account, asset, resource
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from datetime import datetime

router = APIRouter(prefix="/export", tags=["导出"])


def _crud_for(module: str):
    table = {
        "villager": villager,
        "household": household,
        "contact": contact,
        "bank_account": bank_account,
        "asset": asset,
        "resource": resource,
    }.get(module)
    if not table:
        raise HTTPException(status_code=400, detail=f"未知模块: {module}")
    return table


def _headers_for(module: str):
    return {
        "villager": ["姓名", "性别", "身份证号", "出生日期", "民族", "文化程度", "职业", "与户主关系", "户号", "住址", "备注", "是否锁定"],
        "household": ["户号", "所属自然村", "户主姓名", "住址", "成员数量", "是否锁定"],
        "contact": ["村民", "类型", "联系方式", "主联系", "备注", "是否锁定"],
        "bank_account": ["村民", "开户行", "账号", "户名", "账号类型", "状态", "备注", "是否锁定"],
        "asset": ["名称", "编码", "类型", "位置", "面积", "数量", "单位", "购置日期", "购置价格", "当前价值", "状态", "备注", "是否锁定"],
        "resource": ["名称", "编码", "类型", "位置", "面积", "储量", "状态", "开发情况", "描述", "备注", "是否锁定"],
    }.get(module, [])


def _household_export_data(db: Session, household_id: int):
    """查询户主姓名、成员数量、所属自然村名称，用于导出"""
    from sqlalchemy import text
    row = db.execute(text("""
        SELECT
            h.household_no,
            nv.name as natural_village_name,
            head.name as head_name,
            h.address,
            COUNT(m.id) as member_count,
            h.is_locked
        FROM households h
        LEFT JOIN natural_villages nv ON h.natural_village_id = nv.id
        LEFT JOIN villagers head ON h.id = head.household_id AND head.relation_to_head = '户主'
        LEFT JOIN villagers m ON h.id = m.household_id
        WHERE h.id = :household_id
        GROUP BY h.id, nv.name, head.name
    """), {"household_id": household_id}).fetchone()
    return row


def _row_from(module: str, obj: any, db: Session):
    is_locked = "是" if getattr(obj, "is_locked", 0) else "否"
    if module == "villager":
        return [
            getattr(obj, "name", ""),
            getattr(obj, "gender", ""),
            getattr(obj, "id_card", ""),
            getattr(obj, "birth_date", ""),
            getattr(obj, "ethnicity", ""),
            getattr(obj, "education", ""),
            getattr(obj, "occupation", ""),
            getattr(obj, "relation_to_head", ""),
            getattr(obj, "household_no", ""),
            getattr(obj, "address", ""),
            getattr(obj, "remark", ""),
            is_locked,
        ]
    elif module == "household":
        # 导出时重新查询 enriched 数据（get_by_id 只返回原始字段）
        extra = _household_export_data(db, obj.id)
        return [
            extra.household_no if extra else "",
            extra.natural_village_name if extra else "",
            extra.head_name if extra else "",
            extra.address if extra else "",
            extra.member_count if extra else 0,
            "是" if (extra.is_locked if extra else 0) else "否",
        ]
    elif module == "contact":
        v = db.query(villager.model).filter(villager.model.id == obj.villager_id).first()
        return [
            v.name if v else "",
            getattr(obj, "type", ""),
            getattr(obj, "value", ""),
            "是" if getattr(obj, "is_primary", 0) == 1 else "否",
            getattr(obj, "remark", ""),
            is_locked,
        ]
    elif module == "bank_account":
        v = db.query(villager.model).filter(villager.model.id == obj.villager_id).first()
        return [
            v.name if v else "",
            getattr(obj, "bank_name", ""),
            getattr(obj, "account_number_encrypted", ""),
            getattr(obj, "account_holder", ""),
            getattr(obj, "account_type", ""),
            "正常" if getattr(obj, "is_active", 1) == 1 else "已注销",
            getattr(obj, "remark", ""),
            is_locked,
        ]
    elif module == "asset":
        return [
            getattr(obj, "name", ""),
            getattr(obj, "code", ""),
            getattr(obj, "asset_type", ""),
            getattr(obj, "location", ""),
            getattr(obj, "area", ""),
            getattr(obj, "quantity", ""),
            getattr(obj, "unit", ""),
            getattr(obj, "purchase_date", ""),
            str(getattr(obj, "purchase_price", "")),
            str(getattr(obj, "current_value", "")),
            getattr(obj, "status", ""),
            getattr(obj, "remark", ""),
            is_locked,
        ]
    elif module == "resource":
        return [
            getattr(obj, "name", ""),
            getattr(obj, "code", ""),
            getattr(obj, "resource_type", ""),
            getattr(obj, "location", ""),
            getattr(obj, "area", ""),
            getattr(obj, "reserves", ""),
            getattr(obj, "status", ""),
            getattr(obj, "development", ""),
            getattr(obj, "description", ""),
            getattr(obj, "remark", ""),
            is_locked,
        ]
    return [str(getattr(obj, k, "")) for k in dir(obj) if not k.startswith("_")]


@router.post("")
def export_api(body: dict, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    module: str = body.get("module")
    ids: list = body.get("ids", [])
    if not module:
        raise HTTPException(status_code=400, detail="缺少 module 参数")
    if not ids:
        raise HTTPException(status_code=400, detail="请选择要导出的记录")

    crud = _crud_for(module)
    headers = _headers_for(module)

    # 查询数据
    records = []
    for i in ids:
        obj = crud.get_by_id(db, i)
        if obj:
            records.append(obj)

    if not records:
        raise HTTPException(status_code=404, detail="未找到任何记录")

    # 生成 Excel
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = module

    # 表头样式
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    center = Alignment(horizontal="center", vertical="center")
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # 写表头
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center
        cell.border = border

    # 写数据行
    for row_idx, obj in enumerate(records, 2):
        row_data = _row_from(module, obj, db)
        for col_idx, val in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border

    # 调整列宽
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            try:
                max_len = max(max_len, len(str(cell.value or "")))
            except:
                pass
        ws.column_dimensions[col_letter].width = min(max_len + 2, 40)

    # 写入 buffer
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"{module}_{datetime.now().strftime('%Y%m%d%H%M%S')}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )
